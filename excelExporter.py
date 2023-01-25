import openpyxl
from openpyxl import load_workbook
from openpyxl import Workbook
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import glob
import os
    
def excelfilecreator(finalsheetsdict, cwstypelist, formattype, compilermode):
    # (dict of dataframes, list of sheet names, target destination)

    newfolder = r'./Compiled/'+f'{formattype}'
    if not os.path.exists(newfolder):
        os.makedirs(newfolder)

    if compilermode == 'standard':
        for sheetname in cwstypelist:
            writer = pd.ExcelWriter('./Compiled/'+f'{formattype}/'+sheetname+f'_{formattype}'+'_Compiled.xlsx', engine='xlsxwriter')
            finalsheetsdict[sheetname].to_excel(writer, sheet_name=sheetname, startrow=1, header=False, index=False)
            workbook = writer.book
            worksheet = writer.sheets[sheetname]
            (max_row, max_col) = finalsheetsdict[sheetname].shape
            column_settings = [{'header': column} for column in finalsheetsdict[sheetname].columns]
            worksheet.add_table(0, 0, max_row, max_col - 1, {'columns': column_settings})
            worksheet.set_column(0, max_col - 1, 12)
            writer.close() 
            
    elif compilermode == 'master':
        writer = pd.ExcelWriter('./Compiled/' + f'{formattype}/' + 'Master' + f'_{formattype}' + '_Compiled.xlsx', engine='xlsxwriter')
        for sheetname in cwstypelist:
            finalsheetsdict[sheetname].to_excel(writer, sheet_name=sheetname, startrow=1, header=False, index=False)
            workbook = writer.book
            worksheet = writer.sheets[sheetname]
            (max_row, max_col) = finalsheetsdict[sheetname].shape
            column_settings = [{'header': column} for column in finalsheetsdict[sheetname].columns]
            worksheet.add_table(0, 0, max_row, max_col - 1, {'columns': column_settings})
            worksheet.set_column(0, max_col - 1, 12)
        writer.close() 
        
    elif compilermode == 'polymerise':
        writer = pd.ExcelWriter('./Compiled/' + f'{formattype}/' + 'Grand Master' + f'_{formattype}' + '_Compiled.xlsx',engine='xlsxwriter')
        polymeriseddict = pd.concat(list(finalsheetsdict.values()))
        polymeriseddict.to_excel(writer, sheet_name='Grand_Master', startrow=1, header=False, index=False)
        workbook = writer.book
        worksheet = writer.sheets['Grand_Master']
        (max_row, max_col) = polymeriseddict.shape
        column_settings = [{'header': column} for column in polymeriseddict.columns]
        worksheet.add_table(0, 0, max_row, max_col - 1, {'columns': column_settings})
        worksheet.set_column(0, max_col - 1, 12)
        writer.close() 